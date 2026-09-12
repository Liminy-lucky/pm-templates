"""
generate_xlsx.py

Reads MASTER.csv, MILESTONES.csv, PEOPLE.csv, RISK_LOG.csv and generates pm-templates.xlsx with:
- Master, Milestones, People, Risk Log sheets (data)
- Calculated columns: NextMilestoneDate, NextDays, Alert
- Conditional formatting: overdue (red), <=20 days (orange), <=30 days (yellow), external customer highlight
- Simple Dashboard sheet with summary KPIs and a bar chart for person workload

This script is meant to be run in the GitHub Action defined in .github/workflows/build_xlsx.yml
"""

import pandas as pd
from datetime import datetime, date
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, Reference

TODAY = date.today()

def parse_date(x):
    try:
        if pd.isna(x) or x=="":
            return None
        return pd.to_datetime(x).date()
    except Exception:
        return None

# Read CSVs
master = pd.read_csv('MASTER.csv', dtype=str)
milestones = pd.read_csv('MILESTONES.csv', dtype=str)
people = pd.read_csv('PEOPLE.csv', dtype=str)
risk = pd.read_csv('RISK_LOG.csv', dtype=str)

# Ensure date parsing
for col in ['计划交付时间','实际交付时间','下次里程碑日期']:
    if col in master.columns:
        master[col] = master[col].apply(parse_date)

milestones['里程碑计划日期'] = milestones['里程碑计划日期'].apply(parse_date)
milestones['里程碑实际完成日期'] = milestones['里程碑实际完成日期'].apply(parse_date)

# Compute NextMilestoneDate per project (next planned milestone on or after today)
next_dates = {}
for pid, group in milestones.groupby('项目ID'):
    future_dates = [d for d in group['里程碑计划日期'].tolist() if d is not None and d >= TODAY]
    if future_dates:
        next_dates[pid] = min(future_dates)
    else:
        next_dates[pid] = None

master['NextMilestoneDate'] = master['项目ID'].map(next_dates)

# NextDays
def days_until(d):
    if d is None:
        return None
    return (d - TODAY).days

master['NextDays'] = master['NextMilestoneDate'].apply(days_until)

# Alert
def alert_for_row(row):
    status = row.get('当前状态','')
    nd = row.get('NextDays')
    if status == '已完成':
        return ''
    if nd is None:
        return ''
    if nd < 0:
        return '已逾期'
    if nd <= 20:
        return '临近（≤20天）'
    if nd <= 30:
        return '临近（≤30天）'
    return ''

master['Alert'] = master.apply(alert_for_row, axis=1)

# Compute people current load: count of non-completed projects per person
master['主负责人'] = master['主负责人'].fillna('')
active_mask = master['当前状态'] != '已完成'
loads = master[active_mask].groupby('主负责人').size().to_dict()
people['CurrentLoad'] = people['负责人'].map(lambda x: int(loads.get(x,0)))

# Prepare writer
out_file = 'pm-templates.xlsx'
with pd.ExcelWriter(out_file, engine='openpyxl') as writer:
    master.to_excel(writer, sheet_name='Master', index=False)
    milestones.to_excel(writer, sheet_name='Milestones', index=False)
    people.to_excel(writer, sheet_name='People', index=False)
    risk.to_excel(writer, sheet_name='Risk Log', index=False)

# Apply conditional formatting and add simple dashboard
wb = load_workbook(out_file)
ws = wb['Master']

# Find column indices by header
headers = {cell.value:cell.column for cell in ws[1]}

# Expected header names exist
nextdays_col = headers.get('NextDays')
alert_col = headers.get('Alert')
external_col = headers.get('是否外部客户')
status_col = headers.get('当前状态')

if nextdays_col:
    # Apply red fill for overdue (NextDays < 0 and status <> 已完成)
    red_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
    orange_fill = PatternFill(start_color='FFD580', end_color='FFD580', fill_type='solid')
    yellow_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')

    # Using formula rules referencing the NextDays cell in the same row
    max_row = ws.max_row
    nd_col_letter = ws.cell(row=1, column=nextdays_col).column_letter
    status_col_letter = ws.cell(row=1, column=status_col).column_letter if status_col else None

    for row in range(2, max_row+1):
        nd_cell = f"${nd_col_letter}${row}"
        status_cell = f"${status_col_letter}${row}" if status_col_letter else None
        # Overdue: NextDays<0 and status<>"已完成"
        if status_col_letter:
            formula_overdue = f"AND({nd_cell}<0,{status_cell}<>'已完成')"
        else:
            formula_overdue = f"{nd_cell}<0"
        ws.conditional_formatting.add(f'A{row}:Z{row}', FormulaRule(formula=[formula_overdue], stopIfTrue=True, fill=red_fill))
        # <=20
        formula_20 = f"AND({nd_cell}>=0,{nd_cell}<=20,{status_cell}<>'已完成')" if status_col_letter else f"AND({nd_cell}>=0,{nd_cell}<=20)"
        ws.conditional_formatting.add(f'A{row}:Z{row}', FormulaRule(formula=[formula_20], stopIfTrue=True, fill=orange_fill))
        # <=30
        formula_30 = f"AND({nd_cell}>20,{nd_cell}<=30,{status_cell}<>'已完成')" if status_col_letter else f"AND({nd_cell}>20,{nd_cell}<=30)"
        ws.conditional_formatting.add(f'A{row}:Z{row}', FormulaRule(formula=[formula_30], stopIfTrue=True, fill=yellow_fill))

# External customer highlight for rows where 是否外部客户 == '是'
if external_col:
    ext_col_letter = ws.cell(row=1, column=external_col).column_letter
    for row in range(2, ws.max_row+1):
        ext_cell = f"${ext_col_letter}${row}"
        formula_ext = f"{ext_cell}='是'"
        wb.active = ws
        ws.conditional_formatting.add(f'A{row}:Z{row}', FormulaRule(formula=[formula_ext], stopIfTrue=True, fill=PatternFill(start_color='FFE6E6', end_color='FFE6E6', fill_type='solid')))

# Create Dashboard sheet with KPIs
if 'Dashboard' in wb.sheetnames:
    dash = wb['Dashboard']
else:
    dash = wb.create_sheet('Dashboard')

# KPIs
total_projects = len(master)
ongoing = len(master[master['当前状态']=='进行中'])
overdue = len(master[master['Alert']=='已逾期'])
external_count = len(master[master['是否外部客户']=='是'])

dash['A1'] = 'KPI'
dash['B1'] = 'Value'
dash['A2'] = 'Total Projects'
dash['B2'] = total_projects
dash['A3'] = 'Ongoing'
dash['B3'] = ongoing
dash['A4'] = 'Overdue'
dash['B4'] = overdue
dash['A5'] = 'External Customers'
dash['B5'] = external_count

# Workload bar chart from People sheet
pws = wb['People']
# Assume '负责人' in col A and 'CurrentLoad' in last column
people_names = []
loads = []
for r in range(2, pws.max_row+1):
    name = pws.cell(row=r, column=1).value
    load = pws.cell(row=r, column=pws.max_column).value
    try:
        load_val = int(load) if load not in (None,'') else 0
    except:
        load_val = 0
    people_names.append(name)
    loads.append(load_val)

# Write a small table in Dashboard for chart source
start_row = 8
dash.cell(row=start_row-1, column=1, value='Owner')
dash.cell(row=start_row-1, column=2, value='Load')
for i, (n, l) in enumerate(zip(people_names, loads)):
    dash.cell(row=start_row+i, column=1, value=n)
    dash.cell(row=start_row+i, column=2, value=l)

chart = BarChart()
chart.type = 'col'
chart.style = 10
chart.title = 'Workload by Owner'
chart.y_axis.title = 'Project Count'
chart.x_axis.title = 'Owner'

data_ref = Reference(dash, min_col=2, min_row=start_row, max_row=start_row+len(loads)-1)
cats_ref = Reference(dash, min_col=1, min_row=start_row, max_row=start_row+len(loads)-1)
chart.add_data(data_ref, titles_from_data=False)
chart.set_categories(cats_ref)
chart.height = 7
chart.width = 14

dash.add_chart(chart, 'D8')

wb.save(out_file)
print('Generated', out_file)
